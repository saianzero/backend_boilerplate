"""
Text extraction from documents (PDF / DOCX / PPTX) + upload validation.

    doc = DocumentUtil()
    text = await doc.extract_text("/tmp/file.pdf", "application/pdf")
    await doc.validate_upload(upload_file, contents)   # raises HTTPException
"""

import logging

import PyPDF2
from docx import Document
from fastapi import HTTPException, UploadFile
from pptx import Presentation

from src.constants import DOCX_FILE_TYPE, MAX_DOCUMENT_UPLOAD_BYTES, PDF_FILE_TYPE

logger = logging.getLogger(__name__)

ALLOWED_DOCUMENT_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt", ".pptx"}


class DocumentUtil:
    async def read_pdf(self, pdf_path: str) -> str:
        try:
            with open(pdf_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                return "".join(page.extract_text() or "" for page in reader.pages).strip()
        except Exception as e:  # noqa: BLE001
            logger.error(f"Error extracting text from PDF: {e}")
            return ""

    async def read_docx(self, docx_path: str) -> str:
        document = Document(docx_path)
        return "\n".join(p.text for p in document.paragraphs)

    async def read_pptx(self, pptx_path: str) -> str:
        presentation = Presentation(pptx_path)
        chunks = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chunks.append(shape.text)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                if notes:
                    chunks.append(notes)
        return "\n".join(chunks).strip()

    async def read_txt(self, path: str) -> str:
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()

    async def extract_text(self, path: str, file_type: str) -> str:
        """Dispatch on MIME type / extension. Returns "" for unsupported types."""
        ft = (file_type or "").lower()
        if ft == PDF_FILE_TYPE or ft.endswith(".pdf") or ft == "pdf":
            return await self.read_pdf(path)
        if ft == DOCX_FILE_TYPE or ft.endswith(".docx") or ft == "docx":
            return await self.read_docx(path)
        if "presentation" in ft or ft.endswith(".pptx"):
            return await self.read_pptx(path)
        if ft.startswith("text/") or ft.endswith(".txt") or ft == "txt":
            return await self.read_txt(path)
        logger.warning(f"Unsupported file type for text extraction: {file_type}")
        return ""

    async def validate_upload(self, file: UploadFile, contents: bytes) -> None:
        """Basic size/extension checks. Extend with magic-byte checks as needed."""
        if len(contents) == 0:
            raise HTTPException(status_code=400, detail=f"File '{file.filename}' is empty")
        if len(contents) > MAX_DOCUMENT_UPLOAD_BYTES:
            limit_mb = MAX_DOCUMENT_UPLOAD_BYTES // (1024 * 1024)
            raise HTTPException(
                status_code=400, detail=f"File '{file.filename}' exceeds {limit_mb}MB limit"
            )
        filename = (file.filename or "").lower()
        if not any(filename.endswith(ext) for ext in ALLOWED_DOCUMENT_EXTENSIONS):
            raise HTTPException(
                status_code=400,
                detail=f"File '{file.filename}' has unsupported format. Allowed: "
                + ", ".join(sorted(e.lstrip(".").upper() for e in ALLOWED_DOCUMENT_EXTENSIONS)),
            )
